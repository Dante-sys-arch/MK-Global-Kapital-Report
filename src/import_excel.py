"""
Import clippings from Excel reports (e.g. the reports delivered to the client)
into the agent's database, and/or measure how many of them the agent finds.

    python src/import_excel.py report1.xlsx report2.xlsx            # show what would be imported
    python src/import_excel.py report1.xlsx --write                 # import (source "manual")
    python src/import_excel.py report1.xlsx --recall                # which reference rows the agent found itself

Reads every sheet with a header row containing "Link" and "Title"/"Titel".
"""
import argparse
import json
import sys
from datetime import datetime
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent))

from openpyxl import load_workbook

from client_config import data_path, load_client
from util import extract_domain, guess_country, iso_date, norm_title, normalize_outlet, normalize_url

ALIASES = {
    "date": ("publication date", "date", "datum"),
    "outlet": ("media outlet", "outlet", "medium"),
    "title": ("title", "titel"),
    "country": ("country", "land"),
    "type": ("type (print / online)", "type", "typ"),
    "link": ("link", "url"),
    "spokesperson": ("spokesperson", "sprecher"),
}


def read_rows(path):
    wb = load_workbook(path, data_only=True)
    rows = []
    for ws in wb.worksheets:
        header_row, cols = None, {}
        for r in range(1, min(ws.max_row, 30) + 1):
            vals = [str(ws.cell(row=r, column=c).value or "").strip().lower() for c in range(1, ws.max_column + 1)]
            if any(v in ALIASES["link"] for v in vals) and any(v in ALIASES["title"] for v in vals):
                header_row = r
                for i, v in enumerate(vals, 1):
                    for key, names in ALIASES.items():
                        if v in names and key not in cols:
                            cols[key] = i
                break
        if not header_row:
            continue
        for r in range(header_row + 1, ws.max_row + 1):
            rec = {k: ws.cell(row=r, column=c).value for k, c in cols.items()}
            if not rec.get("title") and not rec.get("link"):
                continue
            rec["date"] = iso_date(rec.get("date")) if rec.get("date") else ""
            rec = {k: (str(v).strip() if v is not None and k != "date" else v) for k, v in rec.items()}
            rec["_sheet"] = f"{Path(path).name}:{ws.title}"
            rows.append(rec)
    return rows


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("files", nargs="+")
    ap.add_argument("--client", default="mk-global-kapital")
    ap.add_argument("--write", action="store_true")
    ap.add_argument("--recall", action="store_true")
    args = ap.parse_args()

    cfg = load_client(args.client)
    path = data_path(cfg)
    clips = json.loads(path.read_text(encoding="utf-8")) if path.exists() else []
    by_url = {normalize_url(c.get("link", "")): c for c in clips if c.get("link")}
    by_title_dom = {(norm_title(c.get("title", "")), extract_domain(c.get("link", ""))): c for c in clips}

    ref, seen = [], set()
    for f in args.files:
        for row in read_rows(f):
            key = normalize_url(row.get("link") or "") or ("t:" + norm_title(row.get("title", "")) + "|" + (row.get("outlet") or ""))
            if key in seen:
                continue
            seen.add(key)
            ref.append(row)

    if args.recall:
        auto = found = 0
        missed = []
        for row in ref:
            c = by_url.get(normalize_url(row.get("link") or "")) or by_title_dom.get((norm_title(row.get("title", "")), extract_domain(row.get("link") or "")))
            via = [v for v in (c or {}).get("found_via", []) if v]
            if c and (c.get("source") == "auto" or via):
                auto += 1
            else:
                missed.append(row)
            found += bool(c)
        print(f"Referenz: {len(ref)} Clippings | in Datenbank: {found} | vom Agenten selbst gefunden: {auto} ({auto / max(1, len(ref)):.0%})")
        print("\nNicht selbst gefunden:")
        for m in missed:
            print(f"- {m.get('date','')} | {m.get('outlet','')} | {m.get('title','')} | {m.get('link','')}")
        return

    new = []
    for row in ref:
        link = row.get("link") or ""
        if (link and normalize_url(link) in by_url) or (norm_title(row.get("title", "")), extract_domain(link)) in by_title_dom:
            continue
        outlet = normalize_outlet(row.get("outlet") or "", link)
        new.append({
            "date": row.get("date") or "", "outlet": outlet, "title": row.get("title", ""),
            "country": row.get("country") or guess_country(outlet, link), "type": row.get("type") or "Online",
            "tier": 2, "link": link, "people": [p.strip() for p in (row.get("spokesperson") or "").split(",") if p.strip()],
            "added_at": datetime.now().isoformat(timespec="seconds"), "source": "manual",
            "id": f"x-{int(datetime.now().timestamp())}-{len(new)}", "imported_from": row["_sheet"],
        })
    print(f"Referenz: {len(ref)} | schon vorhanden: {len(ref) - len(new)} | neu: {len(new)}")
    for n in new:
        print(f"+ {n['date']} | {n['outlet']} | {n['title'][:70]}")
    if args.write and new:
        clips.extend(new)
        clips.sort(key=lambda x: x.get("date", ""), reverse=True)
        path.write_text(json.dumps(clips, ensure_ascii=False, indent=2), encoding="utf-8")
        print(f"Gespeichert: {path}")


if __name__ == "__main__":
    main()
