"""
Manual input: everything placed in <data dir>/inbox/ (by the dashboard or the GitHub web upload).

    *.json {"kind": "form", ...}   form entry      -> straight into the clippings (source "manual")
    *.json {"kind": "text", ...}   free text / URL -> Claude reads it   -> review list
    *.txt / *.md                   free text       -> Claude reads it   -> review list
    *.png *.jpg *.jpeg *.webp *.gif *.heic  screenshot -> Claude reads it -> review list
    *.pdf                          e-paper page, press review -> Claude reads it -> review list
    *.xlsx *.xlsm *.csv *.tsv      hit list (e.g. PMG)        -> Claude reads it -> review list

Processed files move to inbox/done/. Duplicates (same title, outlet, date and type) are skipped;
online and print count separately.
"""
import base64
import csv
import io
import json
import re
from datetime import datetime
from pathlib import Path

from client_config import data_path
from fetcher import fetch_article
from util import iso_date, norm_title, normalize_url
from validator import extract_clippings

IMAGE_TYPES = {".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".webp": "image/webp", ".gif": "image/gif"}
MAX_IMAGE_BYTES = 4_500_000
ROWS_PER_CHUNK = 250


def inbox_dir(cfg):
    return data_path(cfg).parent / "inbox"


def _medium(value):
    v = (value or "").strip().lower()
    if "e-paper" in v or "epaper" in v or "e paper" in v:
        return "E-Paper"
    if "print" in v or "zeitung" in v or "heft" in v:
        return "Print"
    return "Online"


def _people(names, matcher):
    """Map free-form names to the configured people (by surname)."""
    out = []
    for n in names or []:
        n_l = str(n).lower()
        for p in matcher.people:
            surname = p["name"].split()[-1].lower()
            if surname in n_l and p["name"] not in out:
                out.append(p["name"])
    return out


def _image_block(path):
    data = path.read_bytes()
    media = IMAGE_TYPES.get(path.suffix.lower())
    if path.suffix.lower() in (".heic", ".heif") or len(data) > MAX_IMAGE_BYTES or not media:
        from PIL import Image
        if path.suffix.lower() in (".heic", ".heif"):
            from pillow_heif import register_heif_opener
            register_heif_opener()
        img = Image.open(io.BytesIO(data)).convert("RGB")
        img.thumbnail((2400, 2400))
        buf = io.BytesIO()
        img.save(buf, format="JPEG", quality=85)
        data, media = buf.getvalue(), "image/jpeg"
    return {"type": "image", "source": {"type": "base64", "media_type": media, "data": base64.b64encode(data).decode()}}


def _table_chunks(path):
    rows = []
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                cells = ["" if v is None else str(v).strip() for v in row]
                if any(cells):
                    rows.append("\t".join(cells))
    else:
        text = path.read_text(encoding="utf-8-sig", errors="replace")
        dialect = "\t" if path.suffix.lower() == ".tsv" or text.count("\t") > text.count(";") else ";"
        if dialect == ";" and text.count(",") > text.count(";"):
            dialect = ","
        for row in csv.reader(io.StringIO(text), delimiter=dialect):
            if any(c.strip() for c in row):
                rows.append("\t".join(c.strip() for c in row))
    if not rows:
        return []
    header, body = rows[0], rows[1:] or [rows[0]]
    return [header + "\n" + "\n".join(body[i:i + ROWS_PER_CHUNK]) for i in range(0, len(body), ROWS_PER_CHUNK)]


def _text_blocks(text):
    """Free text; URLs in it are loaded so Claude sees the article itself."""
    blocks = [{"type": "text", "text": "Eingabe:\n" + text[:20000]}]
    for url in re.findall(r"https?://\S+", text)[:10]:
        art = fetch_article(url.rstrip(").,;"))
        if art.get("ok"):
            blocks.append({"type": "text", "text": f"Inhalt von {url}:\nTitel: {art.get('title','')}\n"
                                                   f"Datum laut Seite: {art.get('date','')}\nAutor: {art.get('author','')}\n"
                                                   f"Medium: {art.get('sitename','')}\nText: {art.get('text','')[:6000]}"})
        else:
            blocks.append({"type": "text", "text": f"{url}: Seite nicht lesbar ({art.get('error','')})"})
    return blocks


def _move_done(path):
    done = path.parent / "done"
    done.mkdir(exist_ok=True)
    target = done / path.name
    n = 1
    while target.exists():
        target = done / f"{path.stem}-{n}{path.suffix}"
        n += 1
    path.rename(target)


def process_inbox(agent):
    """Handle every file in the inbox. Returns a list of per-file results for the run report."""
    folder = inbox_dir(agent.cfg)
    results = []
    if not folder.exists():
        if not agent.dry_run:
            folder.mkdir(parents=True, exist_ok=True)
            (folder / "README.md").write_text(
                "Dateien hier ablegen (Screenshot, PDF, Excel/CSV-Trefferliste, Text). Jeder Upload startet einen Lauf.\n"
                "Formular-Einträge aus dem Dashboard landen hier als .json und gehen direkt in die Clippings.\n"
                "Alles, was Claude aus Dateien ausliest, kommt zur Bestätigung in die Prüfliste.\n", encoding="utf-8")
        return results
    files = sorted(p for p in folder.iterdir() if p.is_file() and p.name not in ("README.md", ".gitkeep"))
    for path in files:
        res = {"file": path.name, "added": 0, "review": 0, "duplicates": 0, "error": ""}
        try:
            _handle(agent, path, res)
        except Exception as exc:
            res["error"] = f"{type(exc).__name__}: {str(exc)[:160]}"
        if not res["error"] or "Claude" not in res["error"]:
            if not agent.dry_run:
                _move_done(path)  # Claude errors stay in the inbox and are retried next run
        results.append(res)
    agent.counts["inbox_files"] = len(files)
    return results


def _handle(agent, path, res):
    ext = path.suffix.lower()
    blocks = None
    if ext == ".json":
        data = json.loads(path.read_text(encoding="utf-8"))
        if data.get("kind") == "form":
            _form_entry(agent, data, res, path.name)
            return
        if data.get("kind") == "text":
            blocks = _text_blocks(data.get("text", ""))
        else:
            res["error"] = "unbekanntes JSON-Format"
            return
    elif ext in (".txt", ".md"):
        blocks = _text_blocks(path.read_text(encoding="utf-8", errors="replace"))
    elif ext in IMAGE_TYPES or ext in (".heic", ".heif"):
        blocks = [_image_block(path), {"type": "text", "text": f"Datei: {path.name}"}]
    elif ext == ".pdf":
        blocks = [{"type": "document", "source": {"type": "base64", "media_type": "application/pdf",
                                                  "data": base64.b64encode(path.read_bytes()).decode()}},
                  {"type": "text", "text": f"Datei: {path.name}"}]
    elif ext in (".xlsx", ".xlsm", ".csv", ".tsv"):
        items = []
        for chunk in _table_chunks(path):
            got, err = extract_clippings(agent.validator, [{"type": "text", "text": f"Trefferliste aus {path.name}:\n{chunk}"}])
            if err:
                res["error"] = err
                return
            items.extend(got)
        _to_review(agent, items, res, path.name)
        return
    elif ext == ".xls":
        res["error"] = "altes Excel-Format: bitte als .xlsx speichern"
        return
    else:
        res["error"] = f"Dateityp {ext} wird nicht unterstützt"
        return
    items, err = extract_clippings(agent.validator, blocks)
    if err:
        res["error"] = err
        return
    _to_review(agent, items, res, path.name)


def _manual_key(title, outlet, d, medium):
    return "manual:" + "|".join([norm_title(title), norm_title(outlet), d or "", medium.lower()])


def _form_entry(agent, data, res, fname):
    title = (data.get("title") or "").strip()
    outlet = (data.get("outlet") or "").strip()
    if not title or not outlet:
        res["error"] = "Titel und Medium sind Pflicht"
        return
    d = iso_date(data.get("date"))
    medium = _medium(data.get("medium") or data.get("type"))
    link = (data.get("link") or "").strip()
    if link and medium == "Online" and normalize_url(link) in agent.by_url:
        res["duplicates"] += 1
        return
    item = {"key": _manual_key(title, outlet, d, medium), "link": link, "title": title, "outlet_given": outlet,
            "source": outlet, "feed_date": d, "medium": medium, "country": data.get("country", ""),
            "people": _people(data.get("people"), agent.matcher), "channels": ["manuell"], "verified": "manuell",
            "source_kind": "manual", "note": (data.get("note") or "").strip(), "author": ""}
    clip = agent._accept(item)
    if clip:
        clip["imported_from"] = fname
        res["added"] += 1
    else:
        res["duplicates"] += 1


def _to_review(agent, items, res, fname):
    seen = set()
    for it in items:
        title = (it.get("title") or "").strip()
        outlet = (it.get("outlet") or "").strip()
        d = iso_date(it.get("date"))
        medium = _medium(it.get("medium"))
        link = (it.get("link") or "").strip()
        key = normalize_url(link) if (link and medium == "Online") else _manual_key(title, outlet, d, medium)
        if key in seen:
            continue
        seen.add(key)
        if (link and medium == "Online" and normalize_url(link) in agent.by_url) or \
                agent._manual_duplicate(title, outlet, d, medium):
            res["duplicates"] += 1
            continue
        item = {"key": key, "link": link, "title": title or link, "outlet_given": outlet, "source": outlet,
                "feed_date": d, "medium": medium, "country": it.get("country", ""),
                "people": _people(it.get("people"), agent.matcher), "channels": ["manuell"],
                "snippet": it.get("note", ""), "source_kind": "manual", "source_file": fname}
        before = len(agent.review)
        agent._to_review(item, f"aus {fname} ausgelesen: bitte bestätigen")
        res["review"] += len(agent.review) - before
