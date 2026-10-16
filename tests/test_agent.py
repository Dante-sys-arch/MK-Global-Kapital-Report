"""
Offline tests for the clipping agent. Network calls are replaced by fakes.
Run: python -m unittest discover -s tests -v
"""
import json
import os
import shutil
import sys
import tempfile
import unittest
from datetime import date, timedelta
from pathlib import Path
from unittest import mock

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "src"))

import client_config  # noqa: E402
import util  # noqa: E402
from matcher import Matcher, detect_language  # noqa: E402

CFG = json.loads((ROOT / "clients" / "mk-global-kapital.json").read_text(encoding="utf-8"))

DE_TEXT = ("Die Nachfrage nach Krediten ist in der Region hoch und die Banken vergeben nicht genug Mittel "
           "an kleine Firmen, die auf Finanzierung angewiesen sind. Das ist auch für Anleger von Bedeutung, "
           "weil die Renditen mit der Region wachsen und sich die Risiken auf viele Kredite verteilen. ")


class FakeResp:
    def __init__(self, status=200, text="", content=None, url="", headers=None, js=None):
        self.status_code = status
        self.text = text
        self.content = content if content is not None else text.encode("utf-8")
        self.url = url
        self.headers = headers or {"Content-Type": "text/html"}
        self._js = js

    def json(self):
        if self._js is None:
            raise ValueError("no json")
        return self._js


class UtilTests(unittest.TestCase):
    def test_http_get_passes_params(self):
        with mock.patch.object(util._session, "get", return_value=FakeResp()) as g:
            util.http_get("https://www.googleapis.com/customsearch/v1", params={"key": "k", "q": "x"})
            self.assertEqual(g.call_args.kwargs["params"], {"key": "k", "q": "x"})

    def test_normalize_url(self):
        a = util.normalize_url("https://www.bondguide.de/topnews/abc/?utm_source=x")
        b = util.normalize_url("http://bondguide.de/topnews/abc")
        self.assertEqual(a, b)

    def test_iso_date(self):
        self.assertEqual(util.iso_date("Tue, 13 Oct 2026 08:00:00 +0200"), "2026-10-13")
        self.assertEqual(util.iso_date("13.04.2026, 10:20"), "2026-04-13")
        self.assertFalse(util.plausible_date("2031-01-01"))


class MatcherTests(unittest.TestCase):
    def setUp(self):
        self.m = Matcher(CFG)

    def test_both_spellings_of_savchenko(self):
        for name in ("Louiza Savchenko", "Louzia Savchenko"):
            r = self.m.analyse("Tokenisierte Anleihen", "Ein Beitrag von " + name + ". " + DE_TEXT, "")
            self.assertEqual(r["decision"], "accept", name)
            self.assertEqual(r["people"], ["Louiza Savchenko"])

    def test_feist_namesake(self):
        r = self.m.analyse("Neuer Markt eröffnet", "Johannes Feist, Marktleiter des Edeka in Gundelfingen, freut sich.", "")
        self.assertEqual(r["decision"], "check")

    def test_feist_finance_context(self):
        r = self.m.analyse("Private Credit in Zentralasien", "Johannes Feist erklärt, wann Anleger profitieren. " + DE_TEXT, "")
        self.assertEqual(r["decision"], "accept")

    def test_brand_with_hyphens(self):
        r = self.m.analyse("x", "Bild: Johannes-Feist-CEO-von-MK-Global-Kapital.jpg", "")
        self.assertTrue(r["brand"])

    def test_byline_from_author_and_text(self):
        self.assertTrue(self.m.analyse("t", DE_TEXT, "Louiza Savchenko")["byline"])
        r = self.m.analyse("t", "Gastbeitrag von Dr. Johannes Feist. " + DE_TEXT + " MK Global Kapital", "")
        self.assertTrue(r["byline"])

    def test_founder_needs_brand(self):
        r = self.m.analyse("Repeople: Wirtschaft wird nachhaltig", "Vincenzo Trani gründete repeople. " + DE_TEXT, "")
        self.assertEqual(r["decision"], "check")
        r = self.m.analyse("x", "Vincenzo Trani, Gründer von Mikro Kapital. " + DE_TEXT, "")
        self.assertEqual(r["decision"], "accept")

    def test_title_suffix(self):
        self.assertEqual(util.clean_title("Wenn Banken passen - DFPA Deutsche Finanz Presse Agentur", "dfpa.info"), "Wenn Banken passen")
        self.assertEqual(util.clean_title('Mikro Kapital: Die "neue Normalität" - zurück zur Realwirtschaft', "Moneycab"),
                         'Mikro Kapital: Die "neue Normalität" - zurück zur Realwirtschaft')

    def test_language(self):
        self.assertEqual(detect_language(DE_TEXT * 2), "de")
        fr = "MK Global Kapital deviendra la société de gestion du fonds et les investisseurs pour la croissance dans les marchés " * 3
        self.assertEqual(detect_language(fr), "fr")


class RssTests(unittest.TestCase):
    def test_description_and_author_are_read(self):
        from sources import Stats, scan_rss
        feed = """<?xml version="1.0"?><rss version="2.0" xmlns:dc="http://purl.org/dc/elements/1.1/"><channel><title>T</title>
        <item><title>Private Credit: Was Anleger jetzt beachten sollten</title><link>https://example.de/a</link>
        <description>Johannes Feist, CEO von MK Global Kapital, erklaert die Risiken.</description></item>
        <item><title>Niedrigwasser und Kreditrisiko</title><link>https://example.de/b</link><dc:creator>Louiza Savchenko</dc:creator>
        <description>Wie Klimarisiken Portfolios veraendern.</description></item>
        <item><title>Anderes Thema</title><link>https://example.de/c</link><description>Nichts.</description></item>
        </channel></rss>"""
        with mock.patch("sources.http_get", return_value=FakeResp(text=feed)):
            out = scan_rss(["https://example.de/feed/"], Matcher(CFG), Stats())
        self.assertEqual(sorted(c["link"] for c in out), ["https://example.de/a", "https://example.de/b"])


class SiteSearchTests(unittest.TestCase):
    def test_wordpress_archive_byline(self):
        from sources import Stats, site_search
        from state import State
        posts = [
            {"link": "https://www.moneycab.com/dossiers/obligation-tokenisiert/", "date": "2026-10-09T17:09:00",
             "title": {"rendered": "Die Obligation ist tokenisiert. Doch womit wird sie bezahlt?"},
             "excerpt": {"rendered": "<p>Tokenisierte Anleihen brauchen digitales Geld.</p>"},
             "content": {"rendered": "<p>Text ... </p><p>Louiza Savchenko, Vorstandsmitglied bei MK Global Kapital</p>"}},
            {"link": "https://www.moneycab.com/x/", "date": "2026-10-01T10:00:00",
             "title": {"rendered": "Feist ist ein seltenes Wort"}, "excerpt": {"rendered": ""}, "content": {"rendered": "<p>nichts</p>"}},
        ]
        st = State.__new__(State)
        st.data = {"checked": {}, "outlets": {}, "daily": {}}
        with mock.patch("sources.http_get", return_value=FakeResp(js=posts)) as g:
            out = site_search(["moneycab.com"], ["Savchenko"], st, Stats(), Matcher(CFG))
        self.assertEqual([c["link"] for c in out], ["https://www.moneycab.com/dossiers/obligation-tokenisiert/"])
        self.assertEqual(out[0]["date"], "2026-10-09")
        self.assertIn("Savchenko", out[0]["snippet"])
        self.assertEqual(g.call_args.kwargs["params"]["search"], '"Savchenko"')
        self.assertTrue(st.data["outlets"]["moneycab.com"]["wp"])


class AgentTests(unittest.TestCase):
    """End-to-end with fake channels and fake article pages."""

    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())
        (self.tmp / "data").mkdir()
        existing = [{
            "date": "2026-03-17", "outlet": "finanznachrichten.de", "id": "m-1", "source": "manual",
            "title": "Private Credit ist nicht gleich Private Credit: Wo die Risiken liegen",
            "link": "https://www.finanznachrichten.de/nachrichten-2026-03/1-private-credit.htm",
            "country": "D", "type": "Online", "tier": 2,
        }]
        (self.tmp / "data" / "clippings.json").write_text(json.dumps(existing), encoding="utf-8")
        self.patches = [
            mock.patch.object(client_config, "REPO_ROOT", self.tmp),
            mock.patch("state.REPO_ROOT", self.tmp),
            mock.patch("runreport.RUNS", self.tmp / "data" / "runs"),
            mock.patch("runreport.REPO_ROOT", self.tmp),
            mock.patch.dict(os.environ, {"GOOGLE_API_KEY": "", "GOOGLE_CX": "", "SERPER_API_KEY": "", "ANTHROPIC_API_KEY": ""}),
        ]
        for p in self.patches:
            p.start()
        import search
        self.search = search

    def tearDown(self):
        for p in self.patches:
            p.stop()
        shutil.rmtree(self.tmp)

    def _page(self, title, body, date_str, author=""):
        meta_author = f'<meta name="author" content="{author}">' if author else ""
        return (f'<html lang="de"><head><title>{title}</title>'
                f'<meta property="article:published_time" content="{date_str}T08:00:00+02:00">{meta_author}'
                f'<meta property="og:site_name" content="Testmedium"></head><body><article><h1>{title}</h1>'
                f'<p>{body}</p><p>{DE_TEXT}</p><p>{DE_TEXT}</p></article></body></html>')

    def test_missed_cases_are_now_found(self):
        pages = {
            # Savchenko byline only (cash-online 13.04.2026); feed date wrong on purpose -> page date wins
            "https://www.cash-online.de/a/tokenisierte-anleihen-715994/": self._page(
                "Tokenisierte Anleihen: Was sich für Anleger ändert", "Louiza Savchenko, MK Global Kapital.", "2026-04-13"),
            # Feist quoted, name not in title (Private Banking Magazin 08.09.2026)
            "https://www.private-banking-magazin.de/private-credit-zentralasien/": self._page(
                "Private Credit in Zentralasien", "Johannes Feist erklärt, wann Anleger davon profitieren.", "2026-09-08"),
            # Syndicated copy of an existing clipping on another outlet -> counts separately
            "https://www.bondguide.de/topnews/private-credit-ist-nicht-gleich/": self._page(
                "Private Credit ist nicht gleich Private Credit: Wo die Risiken liegen", "Ein Kommentar von Johannes Feist, CEO MK Global Kapital.", "2026-03-17"),
            # Edeka namesake -> must not be added unchecked (Claude unavailable -> review list)
            "https://www.badische-zeitung.de/edeka/": self._page(
                "Neuer Markt in Gundelfingen", "Johannes Feist, Marktleiter des Edeka, eröffnet.", "2026-10-01"),
            # French Romandie article -> rejected by rule
            "https://www.allnews.ch/content/mk-global-fr": "<html><body><article><p>" + (
                "MK Global Kapital deviendra la société de gestion du fonds et les investisseurs pour la croissance dans les marchés. " * 4)
                + "</p></article></body></html>",
            # Sitemap article without any MK name -> silently ignored
            "https://www.finews.ch/news/other": self._page("Andere Bank", "Nichts über den Kunden.", "2026-10-15"),
        }

        def fake_get(url, params=None, timeout=15, retries=2, headers=None):
            if url in pages:
                return FakeResp(text=pages[url], url=url)
            return FakeResp(status=404, text="", url=url)

        cands = [
            {"link": "https://www.cash-online.de/a/tokenisierte-anleihen-715994/", "title": "Tokenisierte Anleihen: Was sich für Anleger ändert",
             "snippet": "", "source": "cash-online.de", "date": "2026-01-01", "author": "", "channel": "sitemap", "query": ""},
            {"link": "https://www.private-banking-magazin.de/private-credit-zentralasien/", "title": "", "snippet": "",
             "source": "private-banking-magazin.de", "date": "", "author": "", "channel": "site", "query": ""},
            {"link": "https://www.bondguide.de/topnews/private-credit-ist-nicht-gleich/", "title": "Private Credit ist nicht gleich Private Credit: Wo die Risiken liegen",
             "snippet": "", "source": "bondguide.de", "date": "", "author": "", "channel": "google", "query": ""},
            {"link": "https://www.badische-zeitung.de/edeka/", "title": "Neuer Markt in Gundelfingen", "snippet": "Johannes Feist",
             "source": "badische-zeitung.de", "date": "", "author": "", "channel": "google", "query": ""},
            {"link": "https://www.allnews.ch/content/mk-global-fr", "title": "MK Global Kapital deviendra", "snippet": "MK Global Kapital",
             "source": "allnews.ch", "date": "", "author": "", "channel": "google", "query": ""},
            {"link": "https://www.finews.ch/news/other", "title": "", "snippet": "", "source": "finews.ch", "date": "",
             "author": "", "channel": "sitemap", "query": ""},
        ]
        cfg = client_config._apply_defaults(dict(CFG))
        with mock.patch("fetcher.http_get", side_effect=fake_get), \
                mock.patch.object(self.search.Agent, "collect_daily", return_value=cands), \
                mock.patch.object(self.search.Agent, "syndication_search", return_value=None):
            agent = self.search.Agent(cfg)
            new, report = agent.run("daily")

        by_link = {c["link"]: c for c in new}
        self.assertIn("https://www.cash-online.de/a/tokenisierte-anleihen-715994/", by_link)
        cash = by_link["https://www.cash-online.de/a/tokenisierte-anleihen-715994/"]
        self.assertEqual(cash["date"], "2026-04-13")          # page date, not the wrong feed date
        self.assertEqual(cash["people"], ["Louiza Savchenko"])
        self.assertIn("https://www.private-banking-magazin.de/private-credit-zentralasien/", by_link)
        self.assertEqual(by_link["https://www.private-banking-magazin.de/private-credit-zentralasien/"]["title"], "Private Credit in Zentralasien")
        bond = by_link["https://www.bondguide.de/topnews/private-credit-ist-nicht-gleich/"]
        self.assertEqual(bond["syndication_of"], "m-1")        # syndication kept, linked to original
        self.assertTrue(bond["byline"])
        self.assertNotIn("https://www.badische-zeitung.de/edeka/", by_link)
        self.assertTrue(any(r["link"] == "https://www.badische-zeitung.de/edeka/" for r in agent.review))
        self.assertTrue(any("Romandie" in r["reason"] for r in agent.rejected))
        self.assertEqual(len(new), 3)
        self.assertEqual(report["counts"]["no_match"], 1)

        # Second run: nothing is fetched again, nothing duplicated.
        with mock.patch("fetcher.http_get", side_effect=AssertionError("no refetch")), \
                mock.patch.object(self.search.Agent, "collect_daily", return_value=cands), \
                mock.patch.object(self.search.Agent, "syndication_search", return_value=None):
            agent2 = self.search.Agent(cfg)
            new2, _ = agent2.run("daily")
        self.assertEqual(new2, [])
        saved = json.loads((self.tmp / "data" / "clippings.json").read_text(encoding="utf-8"))
        self.assertEqual(len(saved), 4)

    def test_rolling_backfill_one_month_per_run(self):
        cfg = client_config._apply_defaults(dict(CFG))
        calls = []
        with mock.patch("search.google_news", side_effect=lambda qs, *a, **k: calls.extend(qs) or []):
            agent = self.search.Agent(cfg)
            agent.state.data["backfill_next"] = "2025-07"
            agent._rolling_backfill()
            self.assertEqual(agent.state.data["backfill_next"], "2025-08")
            self.assertTrue(all("after:2025-07-01 before:2025-08-01" in q for q in calls))
            from sources import GNEWS_THROTTLE
            GNEWS_THROTTLE["hit"] = True
            agent._rolling_backfill()
            self.assertEqual(agent.state.data["backfill_next"], "2025-08")  # throttled: no progress
            GNEWS_THROTTLE["hit"] = False

    def test_claude_failure_goes_to_review_not_into_clippings(self):
        from validator import Validator
        from sources import Stats
        v = Validator(dict(CFG), Matcher(CFG), Stats())
        v.client = mock.Mock()
        v.client.messages.create.side_effect = RuntimeError("model retired")
        out = v.decide([{"title": "x", "outlet": "y", "link": "z", "snippet": "", "text": ""}])
        self.assertEqual(out[0][0], "unsure")


class HealthTests(unittest.TestCase):
    def test_no_web_search_is_a_warning(self):
        from runreport import evaluate
        ch = {"google": {"requests": 1, "ok": 0, "errors": 1, "notes": ["Projekt hat keinen Zugang zur Custom Search API"], "candidates": 0},
              "rss": {"requests": 50, "ok": 45, "errors": 5, "notes": [], "candidates": 2}}
        self.assertEqual(evaluate(ch, {}, {})["status"], "warn")

    def test_failing_paid_search_is_an_error(self):
        from runreport import evaluate
        ch = {"serpapi": {"requests": 8, "ok": 0, "errors": 8, "notes": ["HTTP 401 Invalid API key"], "candidates": 0}}
        self.assertEqual(evaluate(ch, {}, {})["status"], "error")

    def test_claude_error_is_an_error(self):
        from runreport import evaluate
        ch = {"claude": {"requests": 1, "ok": 0, "errors": 1, "notes": ["NotFoundError"], "candidates": 0}}
        self.assertEqual(evaluate(ch, {}, {})["status"], "error")


class PlacementTests(unittest.TestCase):
    def test_match_and_overdue(self):
        import placements
        tmp = Path(tempfile.mkdtemp())
        try:
            cfg = {"data_file": "data/clippings.json"}
            with mock.patch("placements.data_path", return_value=tmp / "clippings.json"):
                old = (date.today() - timedelta(days=10)).isoformat()
                (tmp / "placements.csv").write_text(
                    "person;outlet;domain;expected_date;title_hint;status;found_date;found_link;notes\n"
                    f"Johannes Feist;BondGuide;bondguide.de;{old};Niedrigwasser;;;;\n"
                    f"Louiza Savchenko;Moneycab;moneycab.com;{old};;;;;\n", encoding="utf-8")
                clips = [{"link": "https://www.bondguide.de/topnews/niedrigwasser-x/", "title": "Niedrigwasser zeigt die Geografie",
                          "date": old, "people": ["Johannes Feist"]}]
                rep = placements.update_placements(cfg, clips)
            self.assertEqual(len(rep["found_now"]), 1)
            self.assertEqual(len(rep["overdue"]), 1)
            self.assertIn("gefunden", (tmp / "placements.csv").read_text(encoding="utf-8"))
        finally:
            shutil.rmtree(tmp)


class ExcelTests(unittest.TestCase):
    def test_no_tier_column(self):
        import excel_writer
        from openpyxl import load_workbook
        tmp = Path(tempfile.mkdtemp())
        try:
            (tmp / "data").mkdir()
            clips = [{"date": "2026-04-13", "outlet": "cash-online.de", "title": "T", "country": "D", "type": "Online",
                      "tier": 2, "link": "https://x", "people": ["Louiza Savchenko"]}]
            (tmp / "data" / "clippings.json").write_text(json.dumps(clips), encoding="utf-8")
            cfg = client_config._apply_defaults(dict(CFG, latest_copy=None))
            with mock.patch.object(excel_writer, "OUTPUT_DIR", tmp / "out"), \
                    mock.patch.object(client_config, "REPO_ROOT", tmp), \
                    mock.patch("excel_writer.data_path", return_value=tmp / "data" / "clippings.json"):
                path = excel_writer.build_report(cfg)
            ws = load_workbook(path)[cfg["clippings_sheet"]]
            headers = [ws.cell(row=8, column=i).value for i in range(1, 8)]
            self.assertNotIn("Tier", headers)
            self.assertEqual(ws.cell(row=9, column=6).value, "Louiza Savchenko")
        finally:
            shutil.rmtree(tmp)


if __name__ == "__main__":
    unittest.main()
